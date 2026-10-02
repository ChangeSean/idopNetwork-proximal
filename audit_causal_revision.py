"""Check the current manuscript, numeric tables and Word package against records."""
from pathlib import Path
import re
import json
import hashlib
from zipfile import ZipFile
from lxml import etree as E
from docx import Document

ROOT=Path(__file__).resolve().parent
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
def main():
    text=(ROOT/'manuscript.md').read_text(encoding='utf-8')
    checks=[]
    def check(name,condition):
        checks.append((name,bool(condition)))
        assert condition,name
    for n in range(1,8):check(f'section {n}',bool(re.search(rf'^## {n} ',text,re.M)))
    for n in range(1,11):check(f'method 3.{n}',bool(re.search(rf'^### 3\.{n} ',text,re.M)))
    for n in range(1,6):check(f'proposition {n}',f'**Proposition {n} (' in text)
    for n in range(1,5):check(f'theorem {n}',f'**Theorem {n} (' in text)
    for n in ('3a', '3b'):check(f'proposition {n}',f'**Proposition {n} (' in text)
    for n in range(1,6):check(f'appendix A.{n}',f'### A.{n} ' in text)
    check('complete reduced rank algorithm','Bartlett' in text and 'first stage' in text and 'strictly greater than 0.6' in text)
    check('no accidental helper token','oracle_factor###' not in text)
    body,refs=text.split('## References',1)
    numbers=[int(m.group(1)) for m in re.finditer(r'^(\d+)\. ',refs,re.M)]
    seen=[]
    for match in re.finditer(r'\[(\d+(?:,\s*\d+)*)\]',body):
        for n in map(int,match.group(1).split(',')):
            if n not in seen:seen.append(n)
    check('citations first use order',seen==numbers==list(range(1,len(numbers)+1)))
    abstract=text.split('## Abstract',1)[1].split('**Keywords:',1)[0]
    check('abstract under 250 words',len(abstract.split())<=250)
    tables=[]
    for match in re.finditer(r'^\*\*Table ([S\d]+)\.\*\*[^\n]*\n\n((?:\|[^\n]*\n)+)',text,re.M):
        rows=[[cell.strip() for cell in line.strip().strip('|').split('|')] for line in match.group(2).splitlines()]
        tables.append((match.group(1),[rows[0]]+rows[2:]))
    doc=Document(ROOT/'manuscript_SiM.docx')
    check('fifteen tables',len(tables)==len(doc.tables)==15)
    check('main table appearance order',[n for n,_ in tables if n.isdigit()]==list(map(str,range(1,6))))
    numeric_cells=0
    for (number,expected),table in zip(tables,doc.tables):
        check(f'table {number} dimensions',len(table.rows)==len(expected) and len(table.columns)==len(expected[0]))
        if number=='1':continue
        for row_idx,(row,values) in enumerate(zip(table.rows,expected)):
            for col_idx,(cell,value) in enumerate(zip(row.cells,values)):
                check(f'table {number} row {row_idx} cell {col_idx}',cell.text.strip()==value)
                numeric_cells+=1
    with ZipFile(ROOT/'manuscript_SiM.docx') as archive:
        xml=E.fromstring(archive.read('word/document.xml'))
        for table in xml.findall('.//w:tbl',NS):
            check('repeated header',table.find('w:tr/w:trPr/w:tblHeader',NS) is not None)
            check('table width fits page',sum(int(g.get('{'+NS['w']+'}w')) for g in table.findall('w:tblGrid/w:gridCol',NS))<=9072)
        maths=xml.findall('.//m:oMath',NS)
        prose_math=re.sub(r'```.*?```','',text,flags=re.S)
        source_math=len(re.findall(r'\$\$.*?\$\$|(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$)',prose_math,re.S))
        # Five additional native equations are affiliation/author superscripts.
        check('every source equation native in Word',len(maths)==source_math+5)
        plain=' '.join(xml.xpath('//w:t/text()',namespaces=NS))
        check('no unconverted TeX','\\hat' not in plain and '\\operatorname' not in plain and '\\mathrm' not in plain)
        check('all method headings in Word',all(f'3.{n} ' in plain for n in range(1,11)))
    check('Word submission identical',(ROOT/'manuscript_SiM.docx').read_bytes()==(ROOT/'submission/manuscript_SiM.docx').read_bytes())
    for name in ['fig1_schematic','fig2_fits','fig3_causal','fig4_simulation','fig5_cohorts','fig6_application','fig7_clinical_cases','figS1_full_network','figS2_network','figS3_report_rates','figS4_proxy_construction']:
        check(name+' copies',(ROOT/f'figures/{name}.pdf').read_bytes()==(ROOT/f'submission/{name}.pdf').read_bytes())
    result=dict(checks=len(checks),passed=True,numeric_table_cells=numeric_cells,tables=len(tables),references=len(numbers),native_equations=len(maths),
                word_page_rendering='unverified: packaged renderer cannot find soffice.exe',
                source_sha256=hashlib.sha256(text.encode()).hexdigest(),word_sha256=hashlib.sha256((ROOT/'manuscript_SiM.docx').read_bytes()).hexdigest())
    (ROOT/'results/causal_revision_document_audit.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
