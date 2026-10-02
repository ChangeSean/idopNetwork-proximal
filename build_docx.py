"""Build the Statistics in Medicine submission file (Word) from manuscript.md.

    python build_docx.py            -> manuscript_SiM.docx

Layout follows the journal's author guidelines: 12-point Times New Roman, single
spacing, continuous line numbers; title page with full title, short title (<= 70
characters), authors, affiliations and corresponding author; abstract and up to six
keywords; numbered sections; acknowledgements, funding, conflict of interest and data
availability statements before the references; numbered references; figure legends;
tables on separate pages after the reference list; figures supplied as separate files.
"""
import os, re, shutil, subprocess
from docx import Document
from docx.shared import Pt, Inches, Mm, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'manuscript.md'); OUT = os.path.join(HERE, 'manuscript_SiM.docx')
PAGE_BREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'

md = open(SRC, encoding='utf-8').read()
lines = md.split('\n')

# ---- title, abstract, keywords, body, appendix, references, legends, supplementary table
title = lines[0].lstrip('# ').strip()
lines = [l for l in lines[1:] if not l.startswith('*Draft for')]
text = '\n'.join(lines)
def cut(text, start, end):
    i = text.index(start); j = text.index(end, i + len(start)) if end else len(text)
    return text[i:j], text[:i] + text[j:]
abstract_block, text = cut(text, '## Abstract', '## 1 Introduction')
abstract = abstract_block.replace('## Abstract', '').strip()
kw = re.search(r'\*\*Keywords:\*\*\s*(.+)', abstract); keywords = kw.group(1).strip() if kw else ''
abstract = re.sub(r'\*\*Keywords:\*\*.*', '', abstract).replace('---', '').strip()
# order in manuscript.md: body, Software, Figure legends, Appendix A, Supplementary table, References
references, text = cut(text, '## References', None)
supp, text = cut(text, '## Supplementary tables', None) if '## Supplementary tables' in text else ('', text)
appendix, text = cut(text, '## Appendix A', None)
legends, text = cut(text, '## Figure legends', None)
software, text = cut(text, '## Software and data', None)
body = text.strip()

# ---- pull tables (caption + grid) out of the body and the supplement; keep the in-text references
tables = []
def extract_tables(block):
    out = []; L = block.split('\n'); i = 0
    while i < len(L):
        if re.match(r'\*\*Table [S\d]+\.\*\*', L[i]):
            j = i + 1
            while j < len(L) and L[j].strip() == '': j += 1
            k = j
            while k < len(L) and L[k].startswith('|'): k += 1
            tables.append('\n'.join([L[i], ''] + L[j:k])); i = k; continue
        out.append(L[i]); i += 1
    return '\n'.join(out)
body = extract_tables(body); supp = extract_tables(supp)
body = re.sub(r'\n{3,}', '\n\n', body)

# ---- statements (edit before submission)
front = f"""---
title: "{title}"
---

**Short title:** Network-guided proximal causal inference

**Authors:** [First Author]$^{{1}}$, [Second Author]$^{{2}}$, [Corresponding Author]$^{{1,*}}$

$^{{1}}$ [Department, Institution, City, Country]
$^{{2}}$ [Department, Institution, City, Country]

**\\*Correspondence:** [Name], [postal address]. Email: [address]. Telephone: [number].

**ORCID:** [iDs]

**Word count:** main text approximately {len(body.split())} words.

{PAGE_BREAK}
## Abstract

{abstract}

**Keywords:** {keywords}

{PAGE_BREAK}
"""
back = f"""

## Acknowledgements

[To be completed.]

## Funding

[To be completed.]

## Conflict of interest

[Authors to confirm their conflict-of-interest declaration.]

## Data availability statement

{software.replace('## Software and data', '').strip()}

{appendix.strip()}

{supp.replace('## Supplementary tables', '## Supplementary methods and provenance', 1).replace('### Supplementary methods and study provenance', '').strip()}

{PAGE_BREAK}
{references.strip()}

{PAGE_BREAK}
{legends.strip()}

Figures 1–7 and S1–S4 are supplied as separate vector PDF files. Main text labels use 8 pt or larger type at the 175 mm canvas; dense supplementary network labels use smaller type.

{PAGE_BREAK}
## Tables

""" + (PAGE_BREAK + '\n').join(t for t in tables) + "\n"

full = front + body + back
tmp_md = os.path.join(HERE, '_manuscript_build.md'); open(tmp_md, 'w', encoding='utf-8').write(full)

# ---- reference document: Times New Roman 12 pt throughout
ref = os.path.join(HERE, '_reference.docx')
with open(ref, 'wb') as reference_stream:
    subprocess.run(['pandoc', '--print-default-data-file', 'reference.docx'], check=True, stdout=reference_stream)
d = Document(ref)
for st in d.styles:
    try:
        st.font.name = 'Times New Roman'; st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    except Exception: pass
for name, size, bold in (('Normal', 12, False), ('Body Text', 12, False), ('First Paragraph', 12, False), ('Title', 16, True),
                         ('Heading 1', 13, True), ('Heading 2', 12, True), ('Heading 3', 12, True), ('Compact', 11, False), ('Table', 11, False)):
    try:
        s = d.styles[name]; s.font.size = Pt(size); s.font.bold = bold; s.font.italic = False; s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.line_spacing = 1.0; s.paragraph_format.space_after = Pt(6)
    except KeyError: pass
d.save(ref)

subprocess.run(['pandoc', tmp_md, '-f', 'markdown+tex_math_dollars+raw_attribute', '-t', 'docx', '--shift-heading-level-by=-1', '--reference-doc', ref, '-o', OUT], check=True)

# ---- post-process: margins, line numbers, page numbers
doc = Document(OUT)
for sec in doc.sections:
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.5)
    sectPr = sec._sectPr
    ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous'); ln.set(qn('w:distance'), '360')
    pg = sectPr.find(qn('w:pgMar')); (pg.addnext(ln) if pg is not None else sectPr.append(ln))   # schema order: pgSz, pgMar, lnNumType
    footer = sec.footer; p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph(); p.alignment = 1
    run = p.add_run()
    for tag, txt in (('begin', None), (None, 'PAGE'), ('separate', None), ('end', None)):
        if tag:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), tag); run._r.append(el)
        else:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt; run._r.append(el)

# Deliberate widths preserve readable type for the wide numerical tables.
def column_weights(table):
    headers=[cell.text for cell in table.rows[0].cells]
    if 'Discovery / estimation' in headers:return [.65,1.0,.7,.65,.65,1.0,1.0]
    if len(headers)==5 and 'RMST months and set' in headers:return [.6,.8,.35,1.8,1.8]
    if len(headers)==3:return [1.5,2.4,2.1]
    if 'Survival points (95% CI)' in headers:return [0.65,0.95,0.45,1.85,1.85,0.6]
    if 'Survival points and set' in headers:return [0.65,0.9,0.4,1.8,1.8,0.7]
    if 'Same roles' in headers:return [0.65,0.65,0.65]+[1.0]*(len(headers)-3)
    if 'Patients' in headers:return [0.85,0.75,0.7,0.95,0.95,0.8,0.9]
    if 'Estimator' in headers:return [0.4,1.65,0.65]+[1.05]*(len(headers)-3)
    return [0.85,0.45,0.65]+[1.05]*(len(headers)-3)
text_width = int(doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin)
text_twips = text_width // 635
for table in doc.tables:
    weights=column_weights(table)
    assert len(table.columns) == len(weights), 'Table schema changed; review column widths.'
    table.autofit = False
    widths_twips = [int(text_twips * weight / sum(weights)) for weight in weights]
    widths_twips[-1] = text_twips - sum(widths_twips[:-1])
    widths = [width * 635 for width in widths_twips]
    for col, width in zip(table.columns, widths):
        col.width = width
    pr = table._tbl.tblPr
    table_width = pr.find(qn('w:tblW'))
    table_width.set(qn('w:type'), 'dxa'); table_width.set(qn('w:w'), str(text_twips))
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + side)
        for attr, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + attr), value)
        borders.append(border)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, value in [('top', '65'), ('bottom', '65'), ('left', '60'), ('right', '60')]:
        margin = OxmlElement('w:' + side); margin.set(qn('w:w'), value); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
    pr.append(margins)
    header = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(header)
    for row_index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for col_index, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = width; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F2F2F2'); cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.alignment = 0 if col_index == 0 or len(table.columns) == 3 else 1
                paragraph.paragraph_format.space_before = Pt(0); paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = row_index == 0
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'; run.font.size = Pt(11); run.font.color.rgb = RGBColor(0, 0, 0)
for paragraph in doc.paragraphs:
    if re.match(r'^Table (?:\d+|S\d+)\.', paragraph.text):
        paragraph.paragraph_format.keep_with_next = True
settings = doc.settings.element
update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
doc.save(OUT)
os.remove(tmp_md); os.remove(ref)
submission = os.path.join(HERE, 'submission')
os.makedirs(submission, exist_ok=True)
shutil.copy2(OUT, os.path.join(submission, os.path.basename(OUT)))
for name in ('fig1_schematic', 'fig2_fits', 'fig3_causal', 'fig4_simulation',
             'fig5_cohorts', 'fig6_application', 'fig7_clinical_cases',
             'figS1_full_network', 'figS2_network', 'figS3_report_rates', 'figS4_proxy_construction'):
    src = os.path.join(HERE, 'figures', name + '.pdf')
    if os.path.exists(src):
        shutil.copy2(src, os.path.join(submission, name + '.pdf'))
print('written', OUT)
